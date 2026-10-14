"""Report placeholders (spec D7): {verdict} and one {key} per question, in the
report's language; the built-in placeholders are the ones REPORT_PLACEHOLDERS
lists."""

import pandas as pd
from docx import Document
from pysepal.i18n import set_locale

from deforestation_alerts.project.models import Question, QuestionOption, default_questions
from deforestation_alerts.project.questions import REPORT_PLACEHOLDERS
from deforestation_alerts.scripts import report_builder
from gui.tile.analysis.alert_files import review_placeholders

ROAD = Question(key="road_access", label="Road access",
                options=[QuestionOption(code="yes", label="Yes")])


def test_review_placeholders_label_answers_and_blank_unanswered_ones():
    set_locale("en")
    row = pd.Series({"description": "logging, Roads", "status": "Confirmed"})
    out = review_placeholders([default_questions()[0], ROAD], row, "Confirmed")
    assert out == {"verdict": "True alert", "loss_driver": "Logging, Roads", "road_access": ""}


def test_the_builder_fills_extra_placeholders(tmp_path):
    doc = Document()
    doc.add_paragraph("Driver: {loss_driver} / {verdict}")
    path = tmp_path / "t.docx"
    doc.save(path)
    filled = report_builder._fill_placeholders(Document(path), {"loss_driver": "Logging", "verdict": "True alert"})
    assert filled.paragraphs[0].text == "Driver: Logging / True alert"


def test_the_builder_s_own_placeholders_are_the_reserved_ones():
    row = pd.Series({"admin1": "A", "admin2": "B", "admin3": "C", "alert_date_min": 2026.001,
                     "alert_date_max": 2026.010, "before_img": "b", "after_img": "a",
                     "alert_sources": "GLAD-L", "area_ha": 1.25})
    assert set(report_builder.row_attributes(row)) == set(REPORT_PLACEHOLDERS)
