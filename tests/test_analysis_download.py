"""``build_export``'s download worker writes four artifacts — gpkg, before/after
GeoTIFF, zip and Word report — and all four are USER data, written straight
into the set folder.

These tests drive the real ``_download_data`` with only the two genuinely
external steps stubbed (the geedim image download and the cartopy figure
render) and then read every artifact back out of the set folder.
"""

import base64
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest

from gui.i18n import t
from gui.store.state_manager import AppState
from gui.tile.analysis import export as export_mod
from gui.tile.analysis import alert_files
from deforestation_alerts.scripts import local_files
from gui.tile.analysis.state import AnalysisState
# ``_in`` / ``_build_button`` / ``_alerts_gdf`` live in conftest: the assembled
# panel's tests (tests/test_analysis_sections.py) drive the same export section
# through ``build_controls`` and need them too.
from tests.conftest import _alerts_gdf, _build_button, _in

# Smallest valid JPEG so python-docx's add_picture accepts the stubbed renders.
_JPEG_1PX = base64.b64decode(
    "/9j/4AAQSkZJRgABAQEAYABgAAD/2wBDAAgGBgcGBQgHBwcJCQgKDBQNDAsLDBkSEw8UHRof"
    "Hh0aHBwgJC4nICIsIxwcKDcpLDAxNDQ0Hyc5PTgyPC4zNDL/wAALCAABAAEBAREA/8QAFAAB"
    "AAAAAAAAAAAAAAAAAAAACf/EABQQAQAAAAAAAAAAAAAAAAAAAAD/2gAIAQEAAD8AKp//2Q=="
)

_BEFORE_INFO = ["Sentinel 2", 1700000000, None, "S2_before", "_010124"]
_AFTER_INFO = ["Sentinel 2", 1700000001, None, "S2_after", "_020224"]
_TIF_NAME = "alert_0_s_010124_s_020224.tif"


@pytest.fixture()
def download_ctx(monkeypatch, data_root, tmp_path):
    """Everything ``_download_data`` reads, plus the two external steps stubbed."""
    from deforestation_alerts import gdal_env
    from deforestation_alerts.scripts import report_builder
    from gui.tile.analysis import paths

    scratch = tmp_path / "scratch"
    monkeypatch.setattr(gdal_env, "scratch_dir", lambda: scratch)

    downloads = []

    def _fake_download(image1, image2, image_name, source1, source2, region):
        downloads.append(image_name)
        Path(image_name).write_bytes(b"geotiff-bytes")
        return image_name

    monkeypatch.setattr(paths, "download_both_images", _fake_download)
    monkeypatch.setattr(
        report_builder,
        "plot_tiff_with_overlay",
        lambda tiff_path, output_path, **kw: Path(output_path).write_bytes(_JPEG_1PX),
    )

    app_state = AppState()
    app_state.alerts_gdf.value = _alerts_gdf()
    app_state.actual_alert_id.value = 0

    # What _download_data reads from the shell's SetContext (passed as both
    # app_tile_model and aux_model): the set's folder and the report template.
    folder = data_root / "set"
    folder.mkdir()
    set_context = SimpleNamespace(recipe_folder_path=str(folder), custom_report_template=None)

    widgets = SimpleNamespace(actual_alert_grid=None)
    # The images picked for the alert (their info lists carry the composite's
    # name); the ee.Images themselves stay None, the download is stubbed.
    analysis = AnalysisState()
    analysis.before_info.set(_BEFORE_INFO)
    analysis.after_info.set(_AFTER_INFO)
    # ...and the panes have drawn them: the ee.Images the composite is
    # exported from. Stand-ins, since the geedim download is stubbed, but
    # present -- a build refuses while a lane is still rebuilding.
    analysis.before_image.set(object())
    analysis.after_image.set(object())
    return SimpleNamespace(
        widgets=widgets, analysis=analysis, app_state=app_state, aux_model=set_context,
        app_tile_model=set_context, folder=folder, scratch=scratch, downloads=downloads,
    )


def _run(ctx):
    return export_mod._download_data(
        ctx.widgets, ctx.analysis, ctx.app_state, ctx.aux_model, ctx.app_tile_model
    )


def test_download_writes_every_artifact_into_the_set_folder(download_ctx, data_files):
    ctx = download_ctx
    folder = ctx.folder

    result = _run(ctx)

    assert result == {
        "zip_path": f"{folder}/alert_0.zip",
        "report_path": f"{folder}/alert_0.docx",
    }
    # The four artifacts, and nothing else (no .part, no GDAL sidecar).
    assert data_files(folder) == sorted(
        folder / name for name in ("alert_0.gpkg", _TIF_NAME, "alert_0.zip", "alert_0.docx")
    )
    # The composite was downloaded once, through a .part of the set folder.
    [downloaded] = ctx.downloads
    assert Path(downloaded).parent == folder and Path(downloaded).name.startswith(_TIF_NAME + ".")
    # The zip carries the two members under their plain names, not the
    # scratch path they were built in.
    assert sorted(zipfile.ZipFile(folder / "alert_0.zip").namelist()) == sorted(
        [_TIF_NAME, "alert_0.gpkg"]
    )
    assert (folder / "alert_0.docx").read_bytes()[:2] == b"PK"
    # The scratch folder the gpkg, zip and report were built in is gone.
    assert list(ctx.scratch.iterdir()) == []


def test_download_reuses_the_composite_in_the_set_folder(download_ctx):
    """A second download must not re-export the composite from GEE."""
    ctx = download_ctx

    _run(ctx)
    assert len(ctx.downloads) == 1

    _run(ctx)
    assert len(ctx.downloads) == 1


def test_a_pick_from_another_collection_is_another_composite(download_ctx):
    """The cache above is keyed on the collections too, not only the dates.

    Sentinel 2, Landsat and Planet all have imagery for the same day, so a
    name built from the date suffixes alone collided: switching one side to
    another collection resolved to the file the first pick had already
    downloaded and ``stage_alert_image`` handed that back as a cache hit --
    the reviewer then analysed, exported and reported imagery from the wrong
    collection, with the right dates printed on it.
    """
    ctx = download_ctx

    _run(ctx)

    # Same day, another collection.
    ctx.analysis.before_info.set(
        ["Landsat", _BEFORE_INFO[1], None, "LS_before", _BEFORE_INFO[4]]
    )
    _run(ctx)

    assert len(ctx.downloads) == 2
    assert len({p.name for p in ctx.folder.glob("*.tif")}) == 2


def test_the_same_pick_is_always_the_same_composite():
    """...and the key is stable, or nothing would ever hit that cache."""
    from gui.tile.analysis import paths

    name = paths.alert_image_name("/f", 0, _BEFORE_INFO, _AFTER_INFO)
    assert name == paths.alert_image_name("/f", 0, list(_BEFORE_INFO), list(_AFTER_INFO))
    # The two sides are told apart, not merged: a before/after swap of the
    # same two collections is a different composite.
    landsat_before = ["Landsat", *_BEFORE_INFO[1:]]
    assert paths.alert_image_name("/f", 0, landsat_before, _AFTER_INFO) != name


def test_a_failed_composite_download_leaves_no_composite_or_part_file(download_ctx, monkeypatch):
    from gui.tile.analysis import paths

    ctx = download_ctx

    def _failing(image1, image2, image_name, *args):
        Path(image_name).write_bytes(b"half")
        raise RuntimeError("geedim failed")

    monkeypatch.setattr(paths, "download_both_images", _failing)
    with pytest.raises(RuntimeError):
        _run(ctx)
    assert not (ctx.folder / _TIF_NAME).exists()
    assert not list(ctx.folder.glob("*.part")) and not list(ctx.folder.glob("*.tif*"))
    assert list(ctx.scratch.iterdir()) == []


def test_a_failed_composite_download_leaves_no_gpkg_in_the_set_folder(download_ctx, monkeypatch):
    """Deferred T2: the gpkg was copied into the set folder before the
    composite was downloaded, so a failed download left ``alert_0.gpkg`` behind."""
    from gui.tile.analysis import paths

    ctx = download_ctx

    def _failing(*args):
        raise RuntimeError("geedim failed")

    monkeypatch.setattr(paths, "download_both_images", _failing)
    with pytest.raises(RuntimeError):
        _run(ctx)
    assert list(ctx.folder.iterdir()) == []


def test_download_uses_the_users_custom_report_template(download_ctx, data_root):
    """A custom template is the user's own file, read where it is."""
    ctx = download_ctx
    template = data_root / "my_template.docx"
    template.write_bytes(Path(alert_files.DEFAULT_REPORT_TEMPLATE).read_bytes())
    ctx.aux_model.custom_report_template = str(template)

    result = _run(ctx)

    assert Path(result["report_path"]).read_bytes()[:2] == b"PK"
    assert template.exists()


def test_download_waits_for_both_panes_to_draw_their_picks(download_ctx):
    """Same mixture from the export side: the composite is named after the
    picked pair and built from the two ee.Images, so a lane that has not
    finished rebuilding is a build refused, not a file named for one pick out
    of another's pixels."""
    ctx = download_ctx
    ctx.analysis.before_image.set(None)  # the before pane is still rebuilding

    assert _run(ctx) is None
    assert ctx.downloads == [] and list(ctx.folder.iterdir()) == []


def test_download_is_a_noop_without_a_recipe(download_ctx):
    ctx = download_ctx
    ctx.app_tile_model.recipe_folder_path = ""

    assert _run(ctx) is None


def test_composite_download_drives_the_task_progress_ring(download_ctx, monkeypatch):
    """The before/after composite download reports geedim tile ticks as task
    progress and then hands the ring back to indeterminate for the zip/report;
    a cache hit downloads nothing, so it only resets the ring."""
    from deforestation_alerts.scripts import gee_progress
    from gui.tile.analysis import paths

    ctx = download_ctx
    fake_download = paths.download_both_images

    def _ticking_download(*args, **kwargs):
        callback = gee_progress._local.callback
        callback(1, 2)
        callback(2, 2)
        return fake_download(*args, **kwargs)

    monkeypatch.setattr(paths, "download_both_images", _ticking_download)

    class _Task:
        def __init__(self):
            self.progress = []

        def set_progress(self, value):
            self.progress.append(value)

    first = _Task()
    export_mod._download_data(
        ctx.widgets, ctx.analysis, ctx.app_state, ctx.aux_model, ctx.app_tile_model,
        task=first,
    )
    assert first.progress == [0.5, 1.0, None]

    cached = _Task()
    export_mod._download_data(
        ctx.widgets, ctx.analysis, ctx.app_state, ctx.aux_model, ctx.app_tile_model,
        task=cached,
    )
    assert cached.progress == [None]


def test_a_report_template_gone_from_disk_refuses_before_any_file_is_written(download_ctx):
    ctx = download_ctx
    folder = Path(ctx.app_tile_model.recipe_folder_path)
    ctx.aux_model.custom_report_template = str(folder.parent / "gone.docx")
    before = sorted(p.name for p in folder.iterdir())
    with pytest.raises(alert_files.ReportTemplateMissing):
        _run(ctx)
    assert sorted(p.name for p in folder.iterdir()) == before


def test_the_report_is_built_from_the_projects_template(download_ctx):
    """A marker paragraph in the user's template reaches the report."""
    import docx

    ctx = download_ctx
    folder = Path(ctx.app_tile_model.recipe_folder_path)
    template = folder.parent / "mine.docx"
    doc = docx.Document(alert_files.DEFAULT_REPORT_TEMPLATE)
    doc.add_paragraph("MARKER-6f1c")
    doc.save(template)
    ctx.aux_model.custom_report_template = str(template)
    result = _run(ctx)
    text = "\n".join(p.text for p in docx.Document(result["report_path"]).paragraphs)
    assert "MARKER-6f1c" in text


def test_the_download_labels_the_questions_and_fills_their_placeholders(download_ctx):
    """The gpkg row carries current questions' columns as labels and not a
    deleted question's; the report fills {verdict} and one {key} per question."""
    import docx
    import geopandas as gpd

    from deforestation_alerts.project.models import default_questions
    from gui.scripts import questions as questions_mod

    ctx = download_ctx
    template = ctx.folder.parent / "mine.docx"
    doc = docx.Document(alert_files.DEFAULT_REPORT_TEMPLATE)
    doc.add_paragraph("PH {verdict} / {loss_driver}")
    doc.save(template)
    ctx.aux_model.custom_report_template = str(template)
    gdf = ctx.app_state.alerts_gdf.value.copy()
    gdf["description"] = "logging, roads"
    gdf["status"] = "Confirmed"
    gdf["q_gone"] = "x"
    ctx.app_state.alerts_gdf.value = gdf
    alert = export_mod.capture_alert(ctx.analysis, ctx.widgets, ctx.app_state)
    alert.questions = default_questions()

    result = export_mod._download_data(
        ctx.widgets, ctx.analysis, ctx.app_state, ctx.aux_model, ctx.app_tile_model, alert=alert
    )

    row = gpd.read_file(ctx.folder / "alert_0.gpkg").iloc[0]
    assert row["description"] == questions_mod.label(default_questions()[0], "logging, roads")
    assert "q_gone" not in gpd.read_file(ctx.folder / "alert_0.gpkg").columns
    text = "\n".join(p.text for p in docx.Document(result["report_path"]).paragraphs)
    assert f"PH True alert / {row['description']}" in text


def test_the_download_writes_the_links_joined_and_no_working_growth_columns(download_ctx):
    """As the main export (spec §6.7): ``grows_from`` joined by ``";"``,
    without ``grows_from_auto`` and ``below_min``."""
    import geopandas as gpd

    ctx = download_ctx
    gdf = ctx.app_state.alerts_gdf.value.copy()
    gdf["grows_from"] = pd.Series([["sep#4@20261002080000000000", "sep#9@20261002080000000000"]]
                                  + [[] for _ in range(len(gdf) - 1)], index=gdf.index, dtype="object")
    gdf["grows_from_auto"] = pd.Series([["sep#4@20261002080000000000"]] + [[] for _ in range(len(gdf) - 1)],
                                       index=gdf.index, dtype="object")
    gdf["below_min"] = False
    ctx.app_state.alerts_gdf.value = gdf

    _run(ctx)

    back = gpd.read_file(ctx.folder / "alert_0.gpkg")
    assert back.iloc[0]["grows_from"] == "sep#4@20261002080000000000;sep#9@20261002080000000000"
    assert "grows_from_auto" not in back.columns and "below_min" not in back.columns


def _grouped(ctx, drawing=None):
    """The context's store as three alerts, the first two reviewed with the
    third into one group (main = alert 1); ``drawing`` is the main's."""
    import geopandas as gpd
    from shapely.geometry import box

    base = ctx.app_state.alerts_gdf.value
    gdf = gpd.GeoDataFrame(pd.concat([base] * 3, ignore_index=True), geometry="geometry", crs="EPSG:4326")
    gdf.geometry = [box(i * 2, 0, i * 2 + 1, 1) for i in range(3)]
    gdf["bounding_box"] = gdf.geometry.envelope
    gdf["alert_polygon"] = gpd.GeoSeries([drawing, None, None], crs="EPSG:4326")
    gdf["status"] = "Confirmed"
    gdf["review_group"] = pd.array([0, 0, 0], dtype="Int64")
    ctx.app_state.alerts_gdf.value = gdf
    return gdf


def test_a_group_download_is_one_feature_with_the_group_outline(download_ctx):
    import geopandas as gpd
    from shapely.geometry import box

    ctx = download_ctx
    drawing = box(0, 0, 5, 1)
    _grouped(ctx, drawing)
    ctx.app_state.actual_alert_id.value = 2  # a member is opened; the files keep its id

    result = _run(ctx)

    assert result["zip_path"] == f"{ctx.folder}/alert_2.zip"
    back = gpd.read_file(ctx.folder / "alert_2.gpkg")
    assert len(back) == 1 and back.crs.to_epsg() == 4326
    assert back.iloc[0]["area_ha"] == pytest.approx(3.75)
    assert back.iloc[0]["group_size"] == 3 and back.iloc[0]["group_alerts"] == "1, 2, 3"
    assert back.geometry.iloc[0].equals(drawing)


def test_a_group_download_without_a_drawing_is_the_union_of_its_alerts(download_ctx):
    import geopandas as gpd

    ctx = download_ctx
    gdf = _grouped(ctx)
    _run(ctx)
    back = gpd.read_file(ctx.folder / "alert_0.gpkg")
    assert back.geometry.iloc[0].equals(gdf.geometry.union_all())


def test_a_single_download_is_unchanged_but_for_two_columns(download_ctx):
    import geopandas as gpd

    ctx = download_ctx
    gdf = ctx.app_state.alerts_gdf.value.copy()
    gdf["grows_from"] = pd.Series([["sep#4@20261002080000000000", "sep#9@20261002080000000000"]],
                                  index=gdf.index, dtype="object")
    ctx.app_state.alerts_gdf.value = gdf

    _run(ctx)

    back = gpd.read_file(ctx.folder / "alert_0.gpkg")
    assert back.iloc[0]["grows_from"] == "sep#4@20261002080000000000;sep#9@20261002080000000000"
    assert back.iloc[0]["group_size"] == 1 and back.iloc[0]["group_alerts"] == "1"
    assert back.iloc[0]["area_ha"] == pytest.approx(1.25)
    assert back.geometry.iloc[0].equals(gdf["bounding_box"].iloc[0])  # no drawing: the box, as before
    assert sorted(set(back.columns) - {"group_size", "group_alerts"}) == sorted(
        set(gpd.GeoDataFrame(gdf).drop(columns=["geometry", "grows_from_auto", "below_min"], errors="ignore").columns)
        - {"bounding_box"} | {"geometry"})


def test_the_report_fills_group_size_and_group_alerts(download_ctx):
    import docx

    ctx = download_ctx
    _grouped(ctx)
    template = ctx.folder.parent / "mine.docx"
    doc = docx.Document(alert_files.DEFAULT_REPORT_TEMPLATE)
    doc.add_paragraph("GROUP {group_size} / {group_alerts}")
    doc.save(template)
    ctx.aux_model.custom_report_template = str(template)

    result = _run(ctx)

    text = "\n".join(p.text for p in docx.Document(result["report_path"]).paragraphs)
    assert "GROUP 3 / 1, 2, 3" in text


# ---------------------------------------------------------------------------
# The export section's list and summary badge (user, 2026-10-26: one list,
# the build as its last row)
# ---------------------------------------------------------------------------
def _project_poke(app_state):
    """A project object to republish ``app_state.project`` with, as
    ``update_project`` does when the report template is picked."""
    from datetime import datetime

    from deforestation_alerts.project.models import Project

    now = datetime.now()
    return Project(name="t", created_at=now, updated_at=now)


def _export_harness(tmp_path, monkeypatch, set_badge=None, defer=False, labels=None):
    """``build_export``'s arguments, wired so a save and a build run inline.

    What the badge says is imperative, per-alert widget state, so it is read
    through the very paths that change it — the shared ``run_save`` and the
    build row's own click handler — with only the blocking bodies stubbed
    and ``spawn_in_context`` running on the calling thread. No notification
    provider is mounted, so the workers publish to a ``NoopNotifier``.

    ``defer=True`` holds each click's worker in ``harness.held`` instead
    (``harness.run()`` runs the oldest), so a test can land a worker *after*
    an alert change — the ordering an inline worker can never produce.
    """
    app_state = AppState()
    app_state.alerts_gdf.value = _alerts_gdf()
    app_state.actual_alert_id.value = 0
    analysis = AnalysisState()
    widgets = SimpleNamespace(actual_alert_grid=None)
    model = SimpleNamespace(recipe_folder_path=str(tmp_path), custom_report_template=None)

    held = []
    monkeypatch.setattr(
        export_mod,
        "spawn_in_context",
        (lambda target, *a, **k: held.append(target)) if defer
        else (lambda target, *a, **k: target()),
    )
    monkeypatch.setattr(export_mod, "_save_attributes", lambda *a, **k: True)
    monkeypatch.setattr(
        export_mod,
        "_download_data",
        lambda *a, **k: {
            "zip_path": str(tmp_path / "alert_0.zip"),
            "report_path": str(tmp_path / "alert_0.docx"),
        },
    )
    # Serve tmp_path, so the built files get a real link.
    monkeypatch.setenv(local_files.DOWNLOAD_ROOT_ENV, str(tmp_path))

    parts = export_mod.build_export(
        widgets=widgets,
        app_state=app_state,
        aux_model=model,
        analysis=analysis,
        app_tile_model=model,
        set_badge=set_badge,
        labels=labels,
    )
    build_tile = _build_button(parts.body)
    assert build_tile is parts.build_tile
    return SimpleNamespace(
        parts=parts,
        analysis=analysis,
        app_state=app_state,
        model=model,
        build_tile=build_tile,
        download_btns=parts.tiles[:2],
        save=parts.run_save,
        build=lambda: build_tile.fire_event("click", None),
        held=held,
        run=lambda: held.pop(0)(),
    )


def test_the_export_badge_follows_what_is_available(tmp_path, monkeypatch):
    seen = []
    harness = _export_harness(
        tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text)
    )
    # The literal English, not ``t()``/``tn()`` of the same key: a missing or
    # misshapen key renders the truthy placeholder ``⟦key⟧`` on both sides of
    # such a comparison, and the assertion passes on nothing.
    assert seen[-1] == "available after save"

    harness.save()  # a save landed: the build is possible
    assert seen[-1] == "ready to build"

    harness.build()  # the zip and the report are there
    assert seen[-1] == "2 files ready"


def _state(harness):
    """What the list shows: the build row's face and whether it can be
    clicked, and whether the zip and the report are live links."""
    zip_tile, docx_tile, _gif = harness.parts.tiles
    row = harness.build_tile
    return SimpleNamespace(
        title=row.title, detail=row.detail, icon=row._glyph.children[0],
        build_enabled=not row.disabled, busy=row.busy,
        files=not zip_tile.disabled and not docx_tile.disabled,
    )


def test_the_export_list_follows_its_states(tmp_path, monkeypatch):
    """User, 2026-10-26: ONE list -- Alert data, Alert report, Timelapse, then
    the build row. BEFORE A SAVE the three file rows and the build row are
    disabled, the build row saying why; SAVED it builds; READY the file rows
    are links and the build row says *Rebuild files*. The Timelapse row is
    there in every state, disabled without a GIF (the alert is not saved
    Confirmed here)."""
    (tmp_path / "alert_0.zip").write_bytes(b"z" * 3072)
    harness = _export_harness(tmp_path, monkeypatch)
    zip_tile, docx_tile, gif_tile = harness.parts.tiles
    assert harness.parts.body.children == [zip_tile, docx_tile, gif_tile, harness.build_tile]
    assert [tile.kind for tile in harness.parts.tiles] == ["zip", "docx", "gif"]

    before = _state(harness)
    assert (before.title, before.detail, before.icon) == ("Build files", "Save this alert first", "mdi-cogs")
    assert (before.build_enabled, before.busy, before.files) == (False, False, False)
    assert gif_tile.disabled

    harness.save()
    saved = _state(harness)
    assert (saved.title, saved.detail, saved.icon) == (
        "Build files", "Makes the files above from the saved review", "mdi-cogs")
    assert (saved.build_enabled, saved.files) == (True, False)
    assert gif_tile.disabled

    harness.build()
    ready = _state(harness)
    assert (ready.title, ready.detail, ready.icon) == (
        "Rebuild files", "Makes them again from the saved review", "mdi-refresh")
    assert (ready.build_enabled, ready.busy, ready.files) == (True, False, True)
    assert gif_tile.disabled
    assert zip_tile.href.endswith("alert_0.zip") and docx_tile.href.endswith("alert_0.docx")
    assert zip_tile.attributes["download"] == "alert_0.zip"
    assert zip_tile.size_text == "3 KB"
    assert docx_tile.size_text == ""  # not on disk in this harness: no size, still linked

    harness.save()  # a new save clears the files: back to *Build files*
    assert _state(harness) == saved
    assert zip_tile.size_text == "" and zip_tile.disabled is True

    harness.build()
    harness.analysis.reset_alert()  # another alert: back to before a save
    assert _state(harness) == before


def test_the_file_rows_stay_disabled_until_the_first_build(tmp_path, monkeypatch):
    """The three file rows are not links before a build has made their files:
    no href, disabled, whatever the save state."""
    harness = _export_harness(tmp_path, monkeypatch)
    for tile in harness.parts.tiles:
        assert tile.disabled is True and tile.path is None
    harness.save()
    for tile in harness.parts.tiles:
        assert tile.disabled is True and tile.path is None
    harness.build()
    assert [tile.disabled for tile in harness.parts.tiles] == [False, False, True]


def test_rebuild_runs_the_build_again(tmp_path, monkeypatch):
    """*Rebuild files* is the same row and the same build; the list stays
    READY with the new files."""
    harness = _export_harness(tmp_path, monkeypatch)
    builds = []
    real = export_mod._download_data
    monkeypatch.setattr(export_mod, "_download_data", lambda *a, **k: builds.append(1) or real(*a, **k))
    harness.save()
    harness.build()
    assert harness.build_tile.title == "Rebuild files"

    harness.build()

    assert builds == [1, 1]
    state = _state(harness)
    assert state.files and state.title == "Rebuild files"
    assert state.build_enabled and not state.busy


def test_the_build_row_spins_while_a_build_runs(tmp_path, monkeypatch):
    """BUILDING: the spinner in place of the icon, the row held and saying
    so, until the worker releases it."""
    harness = _export_harness(tmp_path, monkeypatch, defer=True)
    harness.save()
    harness.run()

    harness.build()

    row = harness.build_tile
    assert (row.title, row.detail) == ("Building files…", "They unlock above when they are ready")
    assert row.busy is True and row.disabled is True
    assert "display:none" in row._glyph.style_ and "display:none" not in row._spinner.style_
    harness.build()  # a click queued while the row was held
    assert len(harness.held) == 1, "a second build started over the first"

    harness.run()

    assert row.busy is False and row.disabled is False and row.title == "Rebuild files"
    assert "display:none" not in row._glyph.style_ and "display:none" in row._spinner.style_


def test_unsaved_changes_hold_the_build(tmp_path, monkeypatch):
    """The files are of the saved review: unsaved work over it holds the
    build row, which says to save the changes; a save gives it back."""
    harness = _export_harness(tmp_path, monkeypatch)
    harness.save()
    assert _state(harness).build_enabled

    harness.analysis.dirty.set(True)
    state = _state(harness)
    assert (state.build_enabled, state.detail) == (False, "Save your changes first")

    harness.analysis.dirty.set(False)
    assert _state(harness).build_enabled


def test_a_template_change_takes_ready_back_to_the_build_tile(tmp_path, monkeypatch):
    harness = _export_harness(tmp_path, monkeypatch)
    harness.save()
    harness.build()
    assert _state(harness).title == "Rebuild files"

    harness.model.custom_report_template = str(tmp_path / "mine.docx")
    harness.app_state.project.value = _project_poke(harness.app_state)

    state = _state(harness)
    assert (state.title, state.build_enabled, state.files) == ("Build files", True, False)


def _dead_links(harness):
    """Whether the two download links offer nothing (no path, both disabled)."""
    return all("alert_0" not in btn.href and btn.disabled is True
               for btn in harness.download_btns)


def test_the_export_badge_does_not_claim_the_previous_alerts_files(tmp_path, monkeypatch):
    """What this section can offer belongs to ONE alert.

    The save stamp leaked across alerts earlier on this branch; the build
    row and the two download links are the same kind of per-alert state.
    An alert change (``reset_alert``) must take the badge back to "available
    after save", and with it the row that would rebuild the previous
    alert's files and the links that still point at them.
    """
    seen = []
    harness = _export_harness(
        tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text)
    )
    harness.save()
    harness.build()
    assert seen[-1] == "2 files ready"

    harness.analysis.reset_alert()  # the reviewer moved on

    assert seen[-1] == "available after save"
    assert harness.build_tile.disabled is True
    assert harness.build_tile.title == "Build files"
    # The links themselves: a later build must not resurrect the previous
    # alert's zip and report.
    assert _dead_links(harness)


def test_a_build_that_lands_after_the_alert_change_offers_nothing(tmp_path, monkeypatch):
    """A worker that finishes after the reviewer moved on applies nothing.

    The build is slow (an imagery download plus the Word report), so the
    reviewer can navigate while it runs. Its apply step is what links the
    file rows: left unguarded, the stale worker offers the previous alert's
    zip and report under a badge counting them, beside the alert now on
    screen.
    """
    seen = []
    harness = _export_harness(
        tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text),
        defer=True,
    )
    harness.save()
    harness.run()  # the save of this alert landed while it was on screen
    assert harness.build_tile.disabled is False

    harness.build()  # the build starts here...
    harness.analysis.reset_alert()  # ...the reviewer moves on...
    harness.run()  # ...and only now does the build land

    assert harness.build_tile.disabled is True  # not re-enabled by the worker
    assert harness.build_tile.busy is False
    assert seen[-1] == "available after save"
    assert _dead_links(harness)


def test_a_build_still_running_does_not_hold_the_next_alerts_row(tmp_path, monkeypatch):
    """The row spins while a build runs, whichever alert is on screen. A save
    of the NEXT alert that lands meanwhile is drawn on that alert's own state
    once the old build ends, so the row is not left disabled over a save."""
    harness = _export_harness(tmp_path, monkeypatch, defer=True)
    harness.save()
    harness.run()
    harness.build()  # the build of this alert starts...
    harness.analysis.reset_alert()  # ...the reviewer moves on...
    harness.save()
    harness.held.pop()()  # ...and saves the next one
    assert harness.build_tile.busy is True  # held while the old build runs

    harness.run()  # the old build lands, stale

    state = _state(harness)
    assert (state.busy, state.build_enabled, state.title, state.files) == (False, True, "Build files", False)


def test_a_later_save_stops_offering_the_files_the_previous_one_built(tmp_path, monkeypatch):
    """The zip and the report are a snapshot of the review that was saved.

    Save -> build -> change the answer -> Save again: the two links and the
    "2 files ready" badge otherwise still hand the reviewer the report built
    from the previous save, for a review that no longer says that. The files
    are not deleted (they stay in the set folder) -- they are simply not this
    review's answer any more, so the list goes back to *Build files*.
    """
    seen = []
    harness = _export_harness(
        tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text)
    )
    harness.save()
    harness.build()
    assert seen[-1] == "2 files ready"

    harness.save()  # the reviewer changed their answer and saved again

    assert seen[-1] == "ready to build"
    assert _dead_links(harness)
    assert harness.build_tile.disabled is False and harness.build_tile.title == "Build files"


def test_other_imagery_stops_offering_the_files_built_from_the_old_pair(tmp_path, monkeypatch):
    """The report and the composite in the zip are of the pair they were built
    from, so another before/after pick invalidates them the same way."""
    seen = []
    harness = _export_harness(
        tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text)
    )
    harness.save()
    harness.build()
    assert seen[-1] == "2 files ready"

    harness.analysis.before_info.set(["Landsat", 1, None, "Landsat 2026-04-10 ", "_100426"])

    assert seen[-1] == "ready to build"
    assert _dead_links(harness)
    assert harness.build_tile.title == "Build files"


def test_another_report_template_stops_offering_the_report_of_the_old_one(tmp_path, monkeypatch):
    """The reviewer's own template is an input of the build too: the ``.docx``
    on the link was written from the previous one."""
    seen = []
    harness = _export_harness(
        tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text)
    )
    harness.save()
    harness.build()
    assert seen[-1] == "2 files ready"

    # What ``ReportTemplateField`` does: write the project, which republishes
    # it; the section reads the template off the live ``SetContext``.
    harness.model.custom_report_template = str(tmp_path / "mine.docx")
    harness.app_state.project.value = _project_poke(harness.app_state)

    assert seen[-1] == "ready to build"
    assert _dead_links(harness)


def test_a_build_that_lands_after_a_later_save_offers_nothing(tmp_path, monkeypatch):
    """The same hole from the other side: the build is in flight when the
    inputs change, and its apply step is what links the files. Without a
    revision captured at the click it restores a readiness that is already
    stale -- the links would point at files built from the review the save has
    just replaced."""
    seen = []
    harness = _export_harness(
        tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text),
        defer=True,
    )
    harness.save()
    harness.run()

    harness.build()        # the build starts here...
    harness.save()
    harness.held.pop()()   # ...a later save of the same alert lands...
    harness.run()          # ...and only now does the build

    assert seen[-1] == "ready to build"
    assert _dead_links(harness)
    assert harness.build_tile.title == "Build files"


def test_a_save_that_lands_after_the_alert_change_offers_no_build(tmp_path, monkeypatch):
    """The save path has the same hole, and the set-wide half still runs.

    A late ``_apply_save`` would enable the build row and badge "ready to
    build" for an alert that was never saved. The ``alerts_gdf`` rebind is NOT
    alert-scoped — the store write did land, and Overview's counts read that
    reactive — so it has to happen even when the alert-scoped half is dropped.
    """
    seen = []
    harness = _export_harness(
        tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text),
        defer=True,
    )
    before = harness.app_state.alerts_gdf.value

    harness.save()  # the save starts here...
    harness.analysis.reset_alert()  # ...the reviewer moves on...
    harness.analysis.dirty.set(True)  # ...and edits the alert now on screen
    harness.run()  # ...and only now does the save land

    assert harness.build_tile.disabled is True
    assert seen[-1] == "available after save"
    # The new alert's unsaved edit is still unsaved: the save wrote the
    # previous alert's review, not this one's.
    assert harness.analysis.dirty.value is True
    # ...but the write did happen, so the set's gdf was still re-published.
    assert harness.app_state.alerts_gdf.value is not before


# ---------------------------------------------------------------------------
# The GIF (spec D9): in the zip, and on its own button
# ---------------------------------------------------------------------------
_GIF_NAME = "alert_0_s_010124_s_020224.gif"


def _gif_job():
    return SimpleNamespace(before_info=_BEFORE_INFO, after_info=_AFTER_INFO, frames=(1, 2, 3))


def _fake_build_gif(job, dest):
    Path(dest).write_bytes(b"GIF89a")
    return Path(dest)


def test_a_confirmed_alerts_zip_carries_its_gif(download_ctx, monkeypatch):
    from gui.tile.analysis import gif as gif_mod

    monkeypatch.setattr(gif_mod, "build_gif", _fake_build_gif)
    download_ctx.app_state.alerts_gdf.value["status"] = "Confirmed"
    result = export_mod._download_data(download_ctx.widgets, download_ctx.analysis, download_ctx.app_state,
                                       download_ctx.aux_model, download_ctx.app_tile_model, gif=_gif_job())
    assert result["gif_path"] == f"{download_ctx.folder}/{_GIF_NAME}"
    assert _GIF_NAME in zipfile.ZipFile(download_ctx.folder / "alert_0.zip").namelist()


def test_no_gif_for_an_alert_not_saved_as_confirmed(download_ctx, monkeypatch):
    from gui.tile.analysis import gif as gif_mod

    called = []
    monkeypatch.setattr(gif_mod, "build_gif", lambda *a: called.append(a))
    download_ctx.app_state.alerts_gdf.value["status"] = "Need revision"
    result = export_mod._download_data(download_ctx.widgets, download_ctx.analysis, download_ctx.app_state,
                                       download_ctx.aux_model, download_ctx.app_tile_model, gif=_gif_job())
    assert called == [] and "gif_path" not in result
    assert _GIF_NAME not in zipfile.ZipFile(download_ctx.folder / "alert_0.zip").namelist()


def test_a_failed_gif_still_builds_the_zip_and_says_so(download_ctx, monkeypatch):
    from gui.tile.analysis import gif as gif_mod

    def _boom(job, dest):
        raise RuntimeError("thumbnail service down")

    monkeypatch.setattr(gif_mod, "build_gif", _boom)
    download_ctx.app_state.alerts_gdf.value["status"] = "Confirmed"
    result = export_mod._download_data(download_ctx.widgets, download_ctx.analysis, download_ctx.app_state,
                                       download_ctx.aux_model, download_ctx.app_tile_model, gif=_gif_job())
    assert result["gif_warning"] is True and "gif_path" not in result
    assert sorted(zipfile.ZipFile(download_ctx.folder / "alert_0.zip").namelist()) == sorted(
        [_TIF_NAME, "alert_0.gpkg"])


def _gif_tile(body):
    """The Timelapse row: disabled until a build makes the GIF, then its
    download link (``path`` set)."""
    from gui.tile.analysis.file_tile import FileTile

    return next(w for w in _in(body) if isinstance(w, FileTile) and w.kind == "gif")


def _gif_harness(tmp_path, monkeypatch, status=None, **kw):
    """``_export_harness`` with the GIF stubbed and the pair picked; its
    build makes a GIF (on disk) along with the zip and the report.

    ``status`` writes the row's verdict and saved pair BEFORE the picks land,
    as for an alert reopened on its saved pair (no save in this session)."""
    from gui.tile.analysis import gif as gif_mod

    harness = _export_harness(tmp_path, monkeypatch, **kw)
    monkeypatch.setattr(gif_mod, "capture_gif", lambda *a, **k: _gif_job())
    monkeypatch.setattr(gif_mod, "build_gif", _fake_build_gif)
    (tmp_path / _GIF_NAME).write_bytes(b"GIF89a")

    def _download(*a, **k):
        paths = {"zip_path": str(tmp_path / "alert_0.zip"), "report_path": str(tmp_path / "alert_0.docx")}
        if k.get("gif") is not None:
            paths["gif_path"] = str(tmp_path / _GIF_NAME)
        return paths

    monkeypatch.setattr(export_mod, "_download_data", _download)
    if status is not None:
        _status(harness, status)
    harness.analysis.before_info.set(_BEFORE_INFO)
    harness.analysis.after_info.set(_AFTER_INFO)
    harness.gif_tile = _gif_tile(harness.parts.body)
    return harness


def _status(harness, value):
    """What a save of the current picks writes into the row: the verdict and
    the saved pair (``info[3]``), which the GIF gate compares the picks with."""
    gdf = harness.app_state.alerts_gdf.value
    gdf["status"] = value
    gdf["before_img"] = _BEFORE_INFO[3]
    gdf["after_img"] = _AFTER_INFO[3]


def _gif_detail(harness):
    return " ".join(_strings(harness.gif_tile._meta))


def _strings(widget):
    out = []
    for c in getattr(widget, "children", None) or []:
        out.extend([c] if isinstance(c, str) else _strings(c))
    return out


def test_the_timelapse_row_waits_for_a_confirmed_save(tmp_path, monkeypatch):
    """Spec D9: only the build makes the GIF, for a true alert on its saved
    pair. Until then the Timelapse row is disabled and says why: nothing
    saved yet, a verdict that is not a true alert, or what the build will
    make."""
    harness = _gif_harness(tmp_path, monkeypatch)
    assert harness.gif_tile.disabled
    assert _gif_detail(harness) == "Save as a true alert, with the images on screen, to create it"
    _status(harness, "Need revision")
    harness.save()
    assert harness.gif_tile.disabled and _gif_detail(harness) == "True alerts only"
    _status(harness, "Confirmed")
    harness.save()
    assert harness.gif_tile.disabled and _gif_detail(harness) == "Animated GIF"
    harness.analysis.reset_alert()  # another alert: its picks are not the saved pair
    assert harness.gif_tile.disabled
    assert _gif_detail(harness) == "Save as a true alert, with the images on screen, to create it"


def test_the_build_of_a_confirmed_alert_links_its_gif_and_counts_it(tmp_path, monkeypatch):
    seen = []
    harness = _gif_harness(tmp_path, monkeypatch, set_badge=lambda text, color=None: seen.append(text))
    _status(harness, "Confirmed")
    harness.save()
    harness.build()
    tile = harness.gif_tile
    assert tile.path is not None and tile.attributes["download"] == _GIF_NAME
    assert not tile.disabled and tile.href and "mdi-download" in _strings(tile)
    assert _gif_detail(harness) == "Animated GIF · 6 B"
    assert seen[-1] == "3 files ready"


def test_a_later_false_positive_save_withdraws_the_gif(tmp_path, monkeypatch):
    """Review Focus #3."""
    harness = _gif_harness(tmp_path, monkeypatch)
    _status(harness, "Confirmed")
    harness.save()
    harness.build()
    assert harness.gif_tile.path is not None
    _status(harness, "False Positive")
    harness.save()
    assert harness.gif_tile.path is None and harness.gif_tile.disabled
    assert _gif_detail(harness) == "True alerts only"


def test_a_zip_build_with_nothing_to_animate_still_offers_the_files_and_says_so(tmp_path, monkeypatch):
    """Spec D9: a GIF that cannot be made costs the GIF, not the build."""
    from gui.tile.analysis import gif as gif_mod

    notifier = _Toasts()
    harness = _gif_harness(tmp_path, monkeypatch)
    monkeypatch.setattr(export_mod, "current_notifier", lambda: notifier)
    passed = {}

    def _refuse(*a, **k):
        raise gif_mod.GifRefused(key="analysis_tile.export_labels.gif_no_images")

    def _download(*a, **k):
        passed["gif"] = k.get("gif")
        return {"zip_path": str(tmp_path / "alert_0.zip"), "report_path": str(tmp_path / "alert_0.docx")}

    monkeypatch.setattr(gif_mod, "capture_gif", _refuse)
    monkeypatch.setattr(export_mod, "_download_data", _download)
    _status(harness, "Confirmed")
    harness.save()
    harness.build()
    assert passed["gif"] is None
    assert not harness.download_btns[0].disabled and harness.gif_tile.path is None
    assert notifier.warnings == [
        "The alert files are ready, but the GIF could not be made, so the zip has no GIF."]


def test_the_timelapse_row_follows_the_saved_pair(tmp_path, monkeypatch):
    """Spec D9: the GIF animates the reviewed pair only. Another pick after
    the Confirmed save means no GIF in the build; picking the saved pair back
    puts it in again."""
    harness = _gif_harness(tmp_path, monkeypatch)
    _status(harness, "Confirmed")
    harness.save()
    assert _gif_detail(harness) == "Animated GIF"
    harness.analysis.before_info.set(["Landsat", 1, None, "LS_other", "_100426"])
    assert _gif_detail(harness) == "Save as a true alert, with the images on screen, to create it"
    harness.analysis.before_info.set(_BEFORE_INFO)
    assert _gif_detail(harness) == "Animated GIF"
    harness.build()
    assert harness.gif_tile.path is not None


def test_a_row_without_a_saved_pair_offers_no_gif(tmp_path, monkeypatch):
    """A missing or empty saved pick is no match: the build makes no GIF."""
    import numpy as np

    harness = _gif_harness(tmp_path, monkeypatch)
    _status(harness, "Confirmed")
    harness.app_state.alerts_gdf.value["after_img"] = np.nan
    harness.save()
    assert _gif_detail(harness) == "Save as a true alert, with the images on screen, to create it"
    harness.build()
    assert harness.gif_tile.path is None and harness.gif_tile.disabled


def test_a_gif_capture_that_crashes_still_builds_the_zip(tmp_path, monkeypatch):
    """Spec D9: a failed GIF does not fail the build -- not even one whose
    capture raises something other than ``GifRefused`` at the click."""
    from gui.tile.analysis import gif as gif_mod

    notifier = _Toasts()
    harness = _gif_harness(tmp_path, monkeypatch)
    monkeypatch.setattr(export_mod, "current_notifier", lambda: notifier)

    def _crash(*a, **k):
        raise RuntimeError("scene index is malformed")

    monkeypatch.setattr(gif_mod, "capture_gif", _crash)
    _status(harness, "Confirmed")
    harness.save()
    harness.build()
    assert not harness.download_btns[0].disabled
    assert not harness.build_tile.disabled and not harness.build_tile.busy
    assert harness.gif_tile.path is None
    assert notifier.warnings == [
        "The alert files are ready, but the GIF could not be made, so the zip has no GIF."]


class _Toasts:
    """A notifier that keeps the warnings and errors it is given."""

    def __new__(cls):
        from pysepal.solara.notifications import NoopNotifier

        class _Notifier(NoopNotifier):
            def __init__(self):
                self.warnings, self.errors = [], []

            def warning(self, message, *, timeout=None):
                self.warnings.append(message)

            def error(self, message, *, timeout=None):
                self.errors.append(message)

        return _Notifier()


def test_a_reopened_true_alert_builds_with_its_gif_without_a_new_save(tmp_path, monkeypatch):
    """Spec D9 (amended), and the build's rule since 2026-10-26: both read
    the row, not this session's saves. An alert saved Confirmed on its pair,
    reopened on that pair, can build at once, and the build makes its GIF."""
    seen = []
    harness = _gif_harness(tmp_path, monkeypatch, status="Confirmed",
                           set_badge=lambda text, color=None: seen.append(text))
    assert not harness.build_tile.disabled
    assert _gif_detail(harness) == "Animated GIF"
    harness.build()
    assert harness.gif_tile.path is not None
    assert harness.gif_tile.attributes["download"] == _GIF_NAME
    assert seen[-1] == "3 files ready"


def test_a_reopened_alert_not_saved_confirmed_offers_no_gif(tmp_path, monkeypatch):
    harness = _gif_harness(tmp_path, monkeypatch, status="Need revision")
    assert not harness.build_tile.disabled  # saved on its pair: it builds
    assert _gif_detail(harness) == "True alerts only"
    harness.build()
    assert harness.gif_tile.path is None and harness.gif_tile.disabled


def test_a_row_the_store_marks_not_reviewed_is_no_saved_review(tmp_path, monkeypatch):
    """The alert store writes "Not reviewed" into every row it has not had a
    save for: that is no saved review, whatever the row's pair. The build
    waits for a save, and the Timelapse row says what to save."""
    harness = _gif_harness(tmp_path, monkeypatch, status="Not reviewed")
    assert harness.build_tile.disabled
    assert harness.build_tile.detail == "Save this alert first"
    assert _gif_detail(harness) == "Save as a true alert, with the images on screen, to create it"
    harness.analysis.dirty.set(True)  # work on an alert never saved: still "save this alert"
    assert harness.build_tile.detail == "Save this alert first"


def test_a_template_change_leaves_the_gif_alone(tmp_path, monkeypatch):
    """Spec D9: the GIF does not depend on the report template, the zip and
    the report do."""
    seen = []
    harness = _gif_harness(tmp_path, monkeypatch, status="Confirmed",
                           set_badge=lambda text, color=None: seen.append(text))
    harness.build()
    harness.model.custom_report_template = str(tmp_path / "mine.docx")
    harness.app_state.project.value = _project_poke(harness.app_state)
    assert harness.gif_tile.path is not None and harness.gif_tile.attributes["download"] == _GIF_NAME
    assert _dead_links(harness)
    assert seen[-1] == "1 file ready"


def test_the_zip_of_a_true_alert_with_an_unsaved_pair_has_no_gif_and_says_so(tmp_path, monkeypatch):
    """Spec D9: the zip follows the gate. A Confirmed row whose pair on
    screen is not the saved pair gets no GIF capture, and one warning."""
    notifier = _Toasts()
    harness = _gif_harness(tmp_path, monkeypatch)
    monkeypatch.setattr(export_mod, "current_notifier", lambda: notifier)
    passed = {}

    def _download(*a, **k):
        passed["gif"] = k.get("gif")
        return {"zip_path": str(tmp_path / "alert_0.zip"), "report_path": str(tmp_path / "alert_0.docx")}

    monkeypatch.setattr(export_mod, "_download_data", _download)
    _status(harness, "Confirmed")
    harness.save()
    harness.analysis.before_info.set(["Landsat", 1, None, "LS_other", "_100426"])  # not saved
    harness.build()
    assert passed["gif"] is None
    assert notifier.warnings == [
        "The alert files are ready, but the GIF could not be made, so the zip has no GIF."]
    assert harness.gif_tile.path is None


def test_a_click_on_the_gif_link_builds_nothing(tmp_path, monkeypatch):
    """Once made, the Timelapse row is the download link: its click is the
    browser's, not a capture or a build -- the row has no click handler at
    all (the standalone GIF button is gone; only the build makes a GIF)."""
    harness = _gif_harness(tmp_path, monkeypatch, status="Confirmed")
    harness.build()
    with pytest.raises(ValueError, match="'click' not found"):
        harness.gif_tile.fire_event("click", None)
    assert harness.gif_tile.path is not None


def test_a_switch_while_a_build_runs_rewords_the_row_and_holds_it(tmp_path, monkeypatch):
    from pysepal.i18n import set_locale

    from gui.tile.analysis.labels import Labels

    labels = Labels()
    harness = _gif_harness(tmp_path, monkeypatch, defer=True, labels=labels)
    _status(harness, "Confirmed")
    harness.save()
    while harness.held:
        harness.run()
    harness.build()
    assert len(harness.held) == 1 and harness.build_tile.busy, "the build worker is held"
    try:
        set_locale("fr")
        labels.run()
        row = harness.build_tile
        assert row.title == t("analysis_tile.export_labels.building") != "Building files…"
        assert row.busy is True and row.disabled is True
        harness.run()                        # the build lands after the switch
        assert row.busy is False and row.title == t("analysis_tile.export_labels.rebuild_tile")
        assert harness.gif_tile.path is not None
    finally:
        set_locale("en")
