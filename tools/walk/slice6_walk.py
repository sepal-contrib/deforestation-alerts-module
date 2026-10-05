"""Slice 6 gate walk (spec §14 row 6): a report from the project's template,
an info button and description on every step, Manage projects' search, the
Project dialog's "saved X ago", and the app in Spanish and French.

    python -m tools.walk.slice6_walk --port 8986 --out /tmp/slice6_walk
"""

import argparse
import json
import re
import sys
import time
from pathlib import Path

import docx
from playwright.sync_api import sync_playwright

from tools.walk import phases as ph
from tools.walk import walklib as w
from tools.walk.seed_project import seed

PROJECT = "slice6_walk"
BOX = (-55.10, -10.10, -55.00, -10.00)  # Amazon, GLAD-L and RADD alerts in 2025-2026
OTHERS = ("north_s6", "north_s6_b", "south_s6")
REPO = Path(__file__).resolve().parents[2]
MESSAGES = REPO / "deforestation_alerts" / "message"
MARKER = "MARKER-SLICE6"
STEP_KEYS = ("aoi", "alerts", "layers", "period", "prepare", "integrate")


def catalogue(locale, name):
    path = MESSAGES / locale / f"{name}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def leaves(node):
    for value in node.values():
        if isinstance(value, dict):
            yield from leaves(value)
        else:
            yield value


def write_template(data: Path) -> Path:
    folder = data / "templates"
    folder.mkdir(parents=True, exist_ok=True)
    doc = docx.Document(str(REPO / "data" / "report_template.docx"))
    doc.add_paragraph(MARKER)
    path = folder / "slice6_template.docx"
    doc.save(str(path))
    return path


def pick_template(page, rec, *, out, data, template):
    """Analysis Export section → the report template field → navigate to the file.

    The section is opened first: the field is a child of the Export body, which
    the review console keeps COLLAPSED (``display: none``), and ``pick_file``
    needs its button ``:visible``. Without this the phase times out and
    ``run_phases`` aborts every phase after it.
    """
    ph.open_section(page, "Export this alert")  # idempotent
    ph.pick_file(page, "Report template", template)  # Task 6's FileInputComponent helper
    w.wait_for(lambda: ph.manifest(data, PROJECT).get("report_template") == str(template), 30,
               "the template saved in project.json")
    rec["report_template"] = ph.manifest(data, PROJECT)["report_template"]
    w.shot(page, out, "template_picked")


def report_has_marker(rec, *, folder):
    reports = sorted(Path(folder).glob("alert_*.docx"), key=lambda p: p.stat().st_mtime)
    assert reports, "no report written"
    text = "\n".join(p.text for p in docx.Document(str(reports[-1])).paragraphs)
    rec["report"] = reports[-1].name
    assert MARKER in text, f"{reports[-1].name} was not built from the project's template"


def info_buttons(page, rec, *, out):
    steps = catalogue("en", "steps")["steps"]
    seen = {}
    for key in STEP_KEYS:
        w.goto_step(page, key, 2000)
        assert steps[key]["description"][:60] in w.panel_text(page), f"{key}: no description"
        page.locator(".right-panel .mdi-information-outline:visible").first.click()
        w.wait_for(lambda: page.locator(ph.MODAL).count(), 15, f"{key}: info dialog")
        text = page.locator(ph.MODAL).last.inner_text()
        assert steps[key]["info_header"] in text, f"{key}: {text[:200]}"
        seen[key] = text[:120]
        page.locator(f'{ph.MODAL} button:has-text("Close")').last.click()
        w.wait_for(lambda: not page.locator(ph.MODAL).count(), 15, f"{key}: info dialog closed")
    ui = catalogue("en", "ui")["analysis_tile"]
    w.goto_step(page, "analysis", 4000)
    w.wait_for(lambda: ui["description"][:60] in w.panel_text(page), 60, "the Analysis panel")

    def _open_info():  # a click while the panel swaps from the last step can land on the leaving tile
        page.locator(".right-panel .mdi-information-outline:visible").last.click()
        page.wait_for_timeout(1500)
        return page.locator(ph.MODAL).count()

    w.wait_for(_open_info, 30, "analysis: info dialog", poll=1)
    assert ui["info_header"] in page.locator(ph.MODAL).last.inner_text()
    # Close, not Escape: the raw dialog never takes the focus, so Escape does
    # not reach it (the first walk left it open over the rail)
    page.locator(f'{ph.MODAL} button:has-text("Close")').last.click()
    w.wait_for(lambda: not page.locator(ph.MODAL).count(), 15, "analysis: info dialog closed")
    rec["info"] = seen
    w.shot(page, out, "info_buttons")


def _listed(page):
    """Which of the seeded projects Manage projects lists now: a line that is
    the name (``north_s6`` must not count on ``north_s6_b``'s line)."""
    text = ph.modal_text(page)
    return {name: bool(re.search(rf"(?m)^\s*{re.escape(name)}(\s|$)", text)) for name in OTHERS}


def project_search(page, rec, *, out):
    """Manage projects' search narrows its list; Cancel leaves the Project
    dialog open for the next phase."""
    w.open_manage(page)
    search = w.en("steps.project.search")
    field = ph.modal(page).locator(f'.v-text-field:has(label:has-text("{search}")) input').first
    w.wait_for(lambda: all(_listed(page).values()), 30, "the seeded projects listed")
    field.fill("north")
    w.wait_for(lambda: _listed(page) == {"north_s6": True, "north_s6_b": True, "south_s6": False}, 15,
               "the search narrows the list")
    field.fill("zz")
    no_match = w.en("steps.project.no_match").format(query="zz")
    w.wait_for(lambda: no_match in ph.modal_text(page) and not any(_listed(page).values()), 15, "the no-match line")
    field.fill("")
    w.wait_for(lambda: all(_listed(page).values()), 15, "the list back")
    rec["listed"] = _listed(page)
    w.shot(page, out, "project_search")
    ph.modal_button(page, w.en("common.cancel")).click()
    w.wait_for(lambda: not ph.dialog_open(page), 15, "Manage projects closed")
    assert w.dialog_visible(page), "Cancel in Manage projects closed the Project dialog"


def _saved_line(page):
    m = re.search(r"Saved (just now|\d+ (?:minute|minutes|hour|hours|day|days) ago)", w.dialog_text(page))
    return m.group(0) if m else None


def saved_ago(page, rec, *, out):
    """The dialog is open (previous phase): its line moves on by itself, then a save resets it."""
    rec["first"] = _saved_line(page)
    assert rec["first"], w.dialog_text(page)[:400]
    t0 = time.time()
    # moved on: a later "N minute(s) ago" than the first line (which may already
    # read "1 minute ago" when the last save is a minute old)
    w.wait_for(lambda: _saved_line(page) != rec["first"] and re.search(r"\d+ minutes? ago", _saved_line(page) or ""),
               150, "the line moved on without reopening the dialog", poll=5)
    rec["later"], rec["waited_s"] = _saved_line(page), round(time.time() - t0)
    w.close_dialog(page)
    w.goto_step(page, "period", 2000)
    lb = page.locator('.right-panel .v-text-field:has(label:has-text("Look-back")) input').first
    t_ms = w.now_ms(page)
    lb.fill("1")
    page.wait_for_timeout(500)
    page.locator('.right-panel button:visible:has-text("Apply")').first.click()  # Task 34's look-back Apply
    w.wait_for(lambda: w.toasts_since(page, t_ms), 30, "the look-back saved (its toast)")
    page.locator('[data-step-id="2"]').first.click()
    w.wait_for(lambda: _saved_line(page) == "Saved just now", 30, "'Saved just now' after the save")
    rec["after_save"] = _saved_line(page)
    page.wait_for_timeout(800)  # the dialog's fade-in, for the screenshot
    w.shot(page, out, "saved_ago")
    w.close_dialog(page)


def _all_leaves(locale) -> set:
    return {v for f in sorted((MESSAGES / "en").glob("*.json")) for v in leaves(catalogue(locale, f.stem))}


def english_left(text, code) -> list:
    """The English catalogue texts (four letters or more, no placeholder, not
    also a text of ``code``'s catalogue) that ``text`` still shows as words."""
    own = _all_leaves(code)
    return sorted(v for v in _all_leaves("en")
                  if "{" not in v and len(re.sub(r"[^A-Za-z]", "", v)) >= 4 and v not in own
                  and re.search(rf"(?<!\w){re.escape(v)}(?!\w)", text))


def _switch(page, code):
    """Pick ``code`` in the language selector and wait until the pipeline
    header reads that language (the prepare segment's tooltip is its label)."""
    page.locator("button:has(.mdi-translate)").first.click()
    page.locator(f'{ph.MODAL} .v-list-item:has-text("({code})")').first.click()
    label = catalogue(code, "app")["app"]["drawer_item"]["prepare"]
    page.locator(f'[data-step-key="prepare"][title^="{label}"]').first.wait_for(state="attached", timeout=60000)


def languages(page, rec, *, out):
    english_legend = {v for v in leaves(catalogue("en", "legend")) if len(re.sub(r"[^A-Za-z]", "", v)) >= 4}
    english_title = catalogue("en", "steps")["steps"]["prepare"]["title"]
    # The active set's layer is switched on in Integrate & vectorize (while the
    # app is still in English) and stays on for the rest of the walk, so every
    # language's status legend can be read off the shared map.
    ph.show_alerts(page)
    for code in ("es-ES", "fr"):
        w.goto_step(page, "prepare", 2000)
        _switch(page, code)
        title = catalogue(code, "steps")["steps"]["prepare"]["title"]
        description = catalogue(code, "steps")["steps"]["prepare"]["description"]
        w.wait_for(lambda: title in w.panel_text(page) and english_title not in w.panel_text(page)
                   and description[:60] in w.panel_text(page), 60, f"{code}: the prepare panel translated")
        body = page.evaluate("() => document.body.innerText")
        assert "⟦" not in body, f"{code}: a missing key renders"
        assert title in w.panel_text(page) and english_title not in w.panel_text(page), f"{code}: step not translated"
        assert description[:60] in w.panel_text(page), f"{code}: the step's description is not translated"
        rec[f"{code}_title"] = title
        for key in STEP_KEYS:  # every step's panel reads translated
            w.goto_step(page, key, 1500)
            left = english_left(w.panel_text(page), code)
            assert not left, f"{code}: the {key} step still reads {left}"
        # Analysis was built in English earlier in the walk. Its build-once
        # layout does not follow a live switch (a known limitation, final
        # review B: accepted for this merge, in docs/sepal-gate.md), so the
        # walk records whether it does instead of failing on it.
        ui = catalogue(code, "ui")["analysis_tile"]
        w.goto_step(page, "analysis", 1500)
        try:
            w.wait_for(lambda: ui["description"][:60] in w.panel_text(page), 15, f"{code}: Analysis translated")
            rec[f"{code}_analysis_follows_switch"] = True
        except AssertionError:
            rec[f"{code}_analysis_follows_switch"] = False
        statuses = list(catalogue(code, "legend")["legend"]["status"].values())
        # Analysis covers the shared map with its own before/after maps and
        # hides the map legend (slice4 asserts it), so the legend is read on a
        # step that shows the shared map.
        w.goto_step(page, "prepare", 1500)
        w.wait_for(lambda: all(v in ph.legend_text(page) for v in statuses), 60,
                   f"{code}: the alert set's legend")
        ph.open_overview(page)  # the dashboard: a dialog, read with dialog_text
        left = english_left(w.dialog_text(page), code)
        assert not left, f"{code}: the Overview still reads {left}"
        legend = ph.legend_text(page)  # LegendComponent's div.sepal-legend (z-index 150; not a v-card)
        rec[f"{code}_legend"] = legend[:300]
        assert legend, f"{code}: no floating legend found"
        leftover = sorted(v for v in english_legend if v in legend and v not in leaves(catalogue(code, "legend")))
        assert not leftover, f"{code}: the legend still reads {leftover}"
        assert all(v in legend for v in statuses), f"{code}: the legend lacks {statuses}"
        w.shot(page, out, f"lang_{code}")
        assert w.close_dialog(page), f"{code}: the Overview dialog did not close"
        page.locator('[data-step-key="prepare"]:visible').first.wait_for(timeout=30000)
    _switch(page, "en")
    w.goto_step(page, "prepare", 2000)
    w.wait_for(lambda: english_title in w.panel_text(page), 60, "the app back in English")
    rec["back_to_en"] = True


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--port", type=int, default=8986)
    ap.add_argument("--out", type=Path, default=Path("/tmp/slice6_walk"))
    ap.add_argument("--data", type=Path, default=Path.home() / "module_results" / "deforestation_alerts_walk_slice6",
                    help="the walk's data dir (DEFORESTATION_ALERTS_DATA_DIR); keep it under $HOME")
    args = ap.parse_args(argv)
    data = args.data
    seed(PROJECT, BOX, sources=("GLAD-L", "RADD"), period=("2025-06-01", "2026-08-31"), look_back=0, root=data)
    for name in OTHERS:
        seed(name, BOX, root=data)
    template = write_template(data)
    args.out.mkdir(parents=True, exist_ok=True)
    proc = w.start_server(args.port, data, args.out / "server.log")
    report, ok = {}, False
    try:
        with sync_playwright() as pw:
            browser, page = w.new_page(pw, report)
            out = args.out
            folder = lambda: ph.set_dir(data, PROJECT)  # noqa: E731
            ok = ph.run_phases(page, [
                ("open", lambda pg, r: (w.open_app(pg, args.port), w.open_project(pg, PROJECT))),
                ("a_prepare", lambda pg, r: ph.prepare_grid_and_run_all(pg, r, out=out)),
                ("a_integrate", lambda pg, r: ph.integrate_and_land(pg, r, out=out)),
                ("b_pick_template", lambda pg, r: pick_template(pg, r, out=out, data=data, template=template)),
                ("b_review_save", lambda pg, r: ph.review_and_save(pg, r, out=out)),
                ("b_pick_images", lambda pg, r: ph.pick_images(pg, r, out=out)),
                ("b_download", lambda pg, r: ph.download_with_images(pg, r, out=out, folder=folder())),
                ("b_report_from_template", lambda pg, r: report_has_marker(r, folder=folder())),
                ("c_info_buttons", lambda pg, r: info_buttons(pg, r, out=out)),
                ("d_project_search", lambda pg, r: project_search(pg, r, out=out)),
                ("e_saved_ago", lambda pg, r: saved_ago(pg, r, out=out)),
                ("f_languages", lambda pg, r: languages(pg, r, out=out)),
            ], report, args.out / "report.json")
            report["error_toasts"] = [x for x in w.toasts_since(page) if x["kind"] == "toast-error"]
            report["page_errors_count"] = len(report.get("page_errors", []))
            browser.close()
        with open(args.out / "report.json", "w") as f:
            json.dump(report, f, indent=1, ensure_ascii=False, default=str)
        ok = ok and not report["error_toasts"] and not report["bad_responses"]
    finally:
        w.stop_server(proc)
    print("WALK", "PASS" if ok else "FAIL", args.out / "report.json")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
